"""由 開源平台佐證文件.md 產生 .docx 與 .pdf（簡易 Markdown：標題、表格、清單、程式碼區塊、段落）。
用法：python gen_doc.py"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted

BASE = Path(__file__).parent
SRC = BASE / "開源平台佐證文件.md"
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def parse(lines):
    """-> list of (kind, payload)"""
    out, i = [], 0
    while i < len(lines):
        ln = lines[i].rstrip("\n")
        if ln.startswith("```"):
            buf = []; i += 1
            while not lines[i].startswith("```"):
                buf.append(lines[i].rstrip("\n")); i += 1
            out.append(("code", "\n".join(buf)))
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            out.append(("table", rows)); continue
        elif ln.startswith("#"):
            n = len(ln) - len(ln.lstrip("#")); out.append((f"h{n}", ln[n:].strip()))
        elif re.match(r"^(- |\d+\. )", ln):
            out.append(("li", re.sub(r"^(- |\d+\. )", "", ln)))
        elif ln.strip() == "---":
            pass
        elif ln.strip():
            out.append(("p", ln.strip()))
        i += 1
    return out


def plain(t):
    return re.sub(r"[`*]", "", t)


def rl_inline(t):
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"`([^`]+)`", r"<font color='#8a1c1c'>\1</font>", t)


def make_docx(items):
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    for k, v in items:
        if k.startswith("h"):
            d.add_heading(plain(v), level=min(int(k[1]), 3) - 1 if k != "h1" else 0)
        elif k == "p":
            d.add_paragraph(plain(v))
        elif k == "li":
            d.add_paragraph(plain(v), style="List Bullet")
        elif k == "code":
            p = d.add_paragraph(); r = p.add_run(v); r.font.name = "Courier New"; r.font.size = Pt(9)
        elif k == "table":
            t = d.add_table(rows=len(v), cols=len(v[0])); t.style = "Table Grid"
            for ri, row in enumerate(v):
                for ci, c in enumerate(row[:len(v[0])]):
                    cell = t.cell(ri, ci); cell.text = plain(c)
                    for r in cell.paragraphs[0].runs:
                        r.font.size = Pt(9); r.bold = ri == 0
            d.add_paragraph()
    d.save(BASE / "開源平台佐證文件.docx")


def make_pdf(items):
    f = "STSong-Light"
    sty = {
        "p": ParagraphStyle("p", fontName=f, fontSize=10, leading=15, spaceAfter=4),
        "h1": ParagraphStyle("h1", fontName=f, fontSize=18, leading=24, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName=f, fontSize=14, leading=20, spaceBefore=10, spaceAfter=6),
        "h3": ParagraphStyle("h3", fontName=f, fontSize=11.5, leading=17, spaceBefore=6, spaceAfter=4),
        "li": ParagraphStyle("li", fontName=f, fontSize=10, leading=15, leftIndent=12, bulletIndent=2),
        "cell": ParagraphStyle("cell", fontName=f, fontSize=8, leading=11),
        "code": ParagraphStyle("code", fontName=f, fontSize=8, leading=11),
    }
    story = []
    for k, v in items:
        if k in ("h1", "h2", "h3", "h4"):
            story.append(Paragraph(rl_inline(v), sty["h3" if k == "h4" else k]))
        elif k == "p":
            story.append(Paragraph(rl_inline(v), sty["p"]))
        elif k == "li":
            story.append(Paragraph(rl_inline(v), sty["li"], bulletText="•"))
        elif k == "code":
            story.append(Preformatted(v, sty["code"]))
        elif k == "table":
            n = len(v[0]); w = (A4[0] - 30 * mm) / n
            data = [[Paragraph(rl_inline(c), sty["cell"]) for c in r[:n]] + [""] * (n - len(r)) for r in v]
            t = Table(data, colWidths=[w] * n, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef5")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 6)]
    SimpleDocTemplate(str(BASE / "開源平台佐證文件.pdf"), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                      topMargin=15 * mm, bottomMargin=15 * mm, title="開源平台佐證文件").build(story)


if __name__ == "__main__":
    items = parse(SRC.read_text(encoding="utf-8").splitlines(True))
    make_docx(items); make_pdf(items)
    print("ok", len(items), "blocks")
